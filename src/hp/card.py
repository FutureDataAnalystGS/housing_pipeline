"""Karta miasta: odpowiedzi na pytania badawcze o lokalną politykę mieszkaniową, każda z DOSŁOWNYMI cytatami.

Pytania są zdefiniowane jawnie (lista QUESTIONS poniżej) — możesz dopisywać własne wzorce. Dla każdej pary miasto × pytanie karta podaje:
  * odpowiedź = jakie statusy wystąpiły (istniejące/realizowane, planowane, zrealizowane, warunkowe, zaprzeczone/wstrzymane, wskazano jako problem),
  * liczbę cytatów, dokumenty źródłowe i do 3 przykładowych cytatów (ze stroną PDF).
Nic nie jest parafrazowane. Brak cytatu = „nie znaleziono w dokumentach” (a nie „miasto tego nie robi”).
"""
from __future__ import annotations

import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path

from . import candidates as C
from .lexicon import score as housing_score
from .reader import _NEG, _VERDICT_CUE


@dataclass
class Q:
    id: str
    grupa: str
    nazwa: str
    rx: re.Pattern
    need_signal: bool = False        # wymagaj czasownika/sygnału działania (ogranicza opisy stanu)
    allow_legal: bool = False        # nie pomijaj zdań z podstawą prawną (art., ustawa)


def _q(id, grupa, nazwa, pattern, **kw):
    return Q(id, grupa, nazwa, re.compile(pattern, re.I), **kw)


QUESTIONS: list[Q] = [
    # ---- działania własne (poziom A)
    _q("A01", "Działania własne", "Budowa mieszkań komunalnych / socjalnych przez gminę",
       r"\bbudow\w+\s+(?:\w+\s+){0,3}(?:mieszka\w+|lokal\w+|budynk\w+)[^.]{0,60}(?:komunaln|socjaln|gminn)|budownictw\w+\s+(?:komunaln|socjaln)|lokali\s+planowanych\s+do\s+wybudowania|powstan\w+\s+\d+\s+(?:mieszka|lokal)|inwestycj\w+\s+mieszkaniow|nowych\s+(?:mieszka|lokal)\w+", need_signal=True),
    _q("A02", "Działania własne", "TBS (Towarzystwo Budownictwa Społecznego) / spółka gminna", r"\bTBS\b|Towarzystw\w+\s+Budownictwa\s+Społeczn|spółk\w+\s+gminn"),
    _q("A03", "Działania własne", "SIM (Społeczna Inicjatywa Mieszkaniowa) / KZN", r"\bSIM\b|Społeczn\w+\s+Inicjatyw\w+\s+Mieszkaniow|\bKZN\b|Krajow\w+\s+Zasob\w+\s+Nieruchomo"),
    _q("A04", "Działania własne", "Zakup, pozyskiwanie lub adaptacja lokali i budynków",
       r"(?:zakup|nabywani|pozyskiwani|nabyci)\w*\s+(?:\w+\s+){0,2}(?:lokal|budynk|mieszka)\w*|lokali\s+planowanych\s+do\s+zakupu|adaptacj\w+\s+(?:budynk|pomieszcze)\w*|przekwalifikowan\w+", need_signal=True),
    _q("A05", "Działania własne", "Najem od prywatnych i podnajem (gmina jako pośrednik)", r"wynajm\w+\s+(?:lokal|mieszka)\w*[^.]{0,40}(?:wolnym\s+rynku|od\s+(?:osób|podmiot))|na\s+wolnym\s+rynku|podnajm\w*"),
    _q("A06", "Działania własne", "Grunty: pozyskiwanie / przeznaczanie pod mieszkalnictwo",
       r"(?:pozyskiw|nabyw|zakup|przeznacz|wydziel)\w*[^.]{0,50}(?:grunt|działk|teren)\w*[^.]{0,60}(?:mieszkani|mieszkalnict|zabudow)|grunt\w*\s+(?:gminn|komunaln)\w*[^.]{0,60}(?:mieszka|zabudow)|wkład\s+(?:gruntow|rzeczow)|bank\w*\s+grunt"),
    _q("A07", "Działania własne", "Grunty i nieruchomości: sprzedaż, zbywanie, przetargi", r"(?:sprzeda|zbyci|zbyw|przetarg|aport)\w*[^.]{0,50}(?:grunt|działk|teren)\w*|(?:grunt|działk)\w*[^.]{0,40}(?:sprzeda|zbyci|przetarg)"),
    _q("A08", "Działania własne", "Wsparcie zakupu mieszkania (dopłaty, kredyty, programy dla młodych)",
       r"Mieszkanie\s+dla\s+Młodych|Rodzinny\s+Kredyt\s+Mieszkaniow|Bezpieczny\s+Kredyt|dopłat\w+\s+do\s+(?:kredyt|zakup|wkładu)|wsparci\w+\s+(?:w\s+|na\s+)?zakup\w*\s+mieszka|program\w*\s+(?:wsparcia|dopłat)[^.]{0,50}zakup|pomoc\w*\s+w\s+zakupie\s+mieszka"),
    _q("A09", "Działania własne", "Preferencje i oferta dla młodych, rodzin, nauczycieli, seniorów",
       r"preferencj\w+[^.]{0,80}(?:młod|rodzin|nauczyciel|senior|najm)|mieszka\w+\s+(?:dla\s+(?:młod|senior|nauczyciel|pracowni|rodzin\s+wielodzietn)|służbow)\w*|młod\w+\s+rodzin\w+|program\w*\s+(?:dla\s+)?senior\w+"),
    _q("A10", "Działania własne", "Mieszkania chronione, wspierane, treningowe", r"mieszka(?:nia|ń|niach|niami|nie|niu)\s+(?:chronion|wspierany|wspieran|treningow)\w*"),
    _q("A11", "Działania własne", "Wychodzenie z bezdomności / usamodzielnienie (Housing First)", r"Housing\s+First|wychodzeni\w+\s+z\s+bezdomności|usamodzielni\w+|pieczy\s+zastępcz\w+"),
    _q("A12", "Działania własne", "Dopłaty do czynszu, „Mieszkanie na start”, Fundusz Dopłat, BGK, KPO",
       r"Mieszkanie\s+na\s+start|dopłat\w+\s+do\s+(?:czynszu|najmu)|Fundusz\w*\s+Dopłat|\bBGK\b|\bKPO\b|Krajow\w+\s+Plan\w*\s+Odbudowy"),
    _q("A13", "Działania własne", "Rewitalizacja obejmująca zasób mieszkaniowy", r"rewitalizac\w+"),
    # ---- zarządzanie zasobem (poziom B)
    _q("B01", "Zarządzanie zasobem", "Sprzedaż lokali (bonifikaty, wykup przez najemców, wstrzymanie sprzedaży)", r"sprzedaż\w*\s+(?:gminnych\s+|komunalnych\s+|wolnych\s+)?(?:lokal|mieszka)\w*|bonifikat\w+|wykup\w*\s+(?:lokal|mieszka)\w*"),
    _q("B02", "Zarządzanie zasobem", "Zamiany lokali, racjonalizacja metrażu", r"zamian\w+\s+(?:lokal|mieszka)\w*|Biur\w+\s+Zamian|nadmetraż\w*|lokal\w+\s+zamienn\w+", need_signal=True),
    _q("B03", "Zarządzanie zasobem", "Remonty i modernizacje zasobu", r"remont\w+|modernizac\w+"),
    _q("B04", "Zarządzanie zasobem", "Termomodernizacja, likwidacja pieców, niska emisja", r"termomodernizac\w*|likwidacj\w+\s+pieców|piec\w+\s+na\s+(?:opał|paliwo)|niskoemisj\w*|niska\s+emisj\w*"),
    _q("B05", "Zarządzanie zasobem", "Polityka czynszowa (stawki, strefy, % wartości odtworzeniowej, weryfikacja dochodów)", r"stawk\w+\s+czynszu|polityk\w+\s+czynszow\w+|weryfikacj\w+\s+(?:dochod|spełnian)\w*|wartoś\w+\s+odtworzeniow\w+|strefy?\s+zamieszkania"),
    _q("B06", "Zarządzanie zasobem", "Windykacja i odpracowanie zadłużenia", r"windykac\w+|egzekucj\w+|odpracowani\w+|rozłożeni\w+\s+na\s+raty|\bugod\w+"),
    _q("B07", "Zarządzanie zasobem", "Zarządca zasobu, zmiany w zarządzaniu, partycypacja najemców", r"zarządc\w+|ZLiBK|\bZBM\b|\bZGL\b|zarządzani\w+\s+(?:zasobem|lokalami)|partycypacyjn\w+|przekształceni\w+\s+(?:zakład|jednostk)"),
    # ---- obowiązki ustawowe (poziom C)
    _q("C01", "Obowiązki ustawowe", "Najem socjalny i wyroki eksmisyjne", r"najem\s+socjaln\w+|lokal\w*\s+socjaln\w+|wyrok\w*\s+(?:o\s+|orzekając\w+\s+)?eksmisj\w+|eksmisj\w+"),
    _q("C02", "Obowiązki ustawowe", "Pomieszczenia tymczasowe, lokale zamienne i kryzysowe", r"pomieszcze\w+\s+tymczasow\w+|lokal\w*\s+zamienn\w+|lokal\w*\s+kryzysow\w+"),
    _q("C03", "Obowiązki ustawowe", "Dodatki mieszkaniowe i obniżki czynszu dla niskich dochodów", r"dodat\w+\s+mieszkaniow\w+|obniżk\w+\s+czynszu|obniżen\w+\s+czynszu"),
    _q("C04", "Obowiązki ustawowe", "Opis obowiązków gminy wynikających z ustawy (cytaty z podstawą prawną)",
       r"obowiązk\w+\s+gminy|obowiązek\s+(?:zapewnieni|dostarczeni|tworzeni)\w*|zadani\w+\s+własn\w+\s+gminy|ochronie\s+praw\s+lokatorów|ustawow\w+\s+obowiązk\w+|art\.\s*4\s+ustawy", allow_legal=True),
    # ---- problemy (diagnoza)
    _q("P01", "Problemy", "Pustostany", r"pustostan\w+|puste\s+lokal\w*|lokal\w*\s+pust\w+|niezasiedlon\w+|długi\s+czas\s+remont\w+"),
    _q("P02", "Problemy", "Zły stan techniczny i wiek zasobu",
       r"zły\s+stan\s+techniczn|stan\s+techniczn\w+[^.]{0,60}(?:niezadowalając|niedostateczn|zły|mierny)|zużyci\w+|wiek\w*\s+budynk|przed\s+19\d\d|najstarsz\w+|katastrof\w+\s+budowlan|nieopłacaln\w+\s+remont|substandard\w+|wymagając\w+\s+(?:remont|znaczn)"),
    _q("P03", "Problemy", "Niedobór mieszkań, kolejki, niezaspokojone potrzeby", r"deficyt\w+|niedobór\w*|brak\w*\s+(?:dostępn\w+\s+)?(?:mieszka|lokali)|oczekując\w+\s+na\s+(?:najem|lokal|mieszka)|zapotrzebowan\w+|kolejk\w+|lista\s+(?:oczekując|osób)|niezaspokojon\w+"),
    _q("P04", "Problemy", "Zaległości w opłatach i zadłużenie najemców", r"zaległo\w+|zadłuż\w+|ściągaln\w+|nieściągnięt\w+"),
    _q("P05", "Problemy", "Odszkodowania (np. za niedostarczenie lokalu)", r"odszkodowan\w+|roszczen\w+"),
    _q("P06", "Problemy", "Bezdomność i wykluczenie mieszkaniowe", r"bezdomn\w+|wykluczeni\w+\s+mieszkaniow"),
    _q("P07", "Problemy", "Niski standard lokali (piece, WC poza lokalem, sutereny, lokale niesamodzielne)", r"WC\s+(?:usytuowane\s+)?poza\s+lokalem|wspóln\w+\s+użytkowani\w+|obniżon\w+\s+standard|niesamodzieln\w+|sutere\w+|piec\w+\s+(?:węglow|na\s+(?:opał|paliwo))"),
    _q("P08", "Problemy", "Niewystarczające środki / koszty wyższe niż wpływy", r"nie\s+pokrywa\w*\s+kosztów|niewystarczając\w+\s+środk|ograniczon\w+\s+środk|brak\w*\s+środków|dopłat\w+\s+z\s+budżetu"),
    # ---- co się udało / ocena
    _q("W01", "Co udało się / ocena", "Zrealizowane działania i wyniki (czas przeszły, dane o wykonaniu)",
       r"\b(?:zrealizow|wybudow|odd|zlikwidow|wyremontow|pozysk|udziel|zakończ|przekaz|zakup|nabyt|sprzedan|utworz|przeprowadz)(?:ono|ano|ała|ał|ały|ili|yły|ane|ona|one|any|ony)\b|osiągnięt\w+|udało\s+się|do\s+końca\s+(?:roku\s+)?20(?:1\d|2[01])"),
    _q("W02", "Co udało się / ocena", "Ocena dotychczasowych działań", r"\bocen\w+\s+(?:realizacji|dotychczas|programu|działań)|dotychczasow\w+\s+(?:realizacj|działa|efekt)|efekt\w+\s+(?:realizacji|programu)|sprawozdani\w+|poprzedni\w+\s+program\w*|monitoring\w*\s+program"),
]
_PAST = QUESTIONS[-2].rx
_HALT = re.compile(r"nie\s+będzie\s+prowadzon\w+|wstrzyman\w+|wstrzymana|zaprzestan\w+", re.I)
CODES = {"istniejące": "I", "planowane": "P", "zrealizowane": "Z", "warunkowe": "W", "zaprzeczone/wstrzymane": "N", "wskazano": "D"}


def status_of(grupa: str, t: str, qid: str = "") -> str:
    if grupa == "Problemy" or qid == "W02":
        return "wskazano"
    if qid == "W01":
        return "zrealizowane"                                 # wzorzec W01 sam jest znacznikiem czasu przeszłego
    if _NEG.search(t) or _HALT.search(t):
        return "zaprzeczone/wstrzymane"
    if grupa != "Co udało się / ocena" and _PAST.search(t):
        return "zrealizowane"
    s = C.act_status(t)
    return {"planned": "planowane", "conditional": "warunkowe"}.get(s, "istniejące")


def answer_questions(doc_id: str, doc_type: str, city: str, pages) -> list[dict]:
    units = C.build_units(pages, max_chars=240 if doc_type == "SR" else 380)
    rows = []
    for i, u in enumerate(units, 1):
        t = u.text
        if doc_type == "SR" and housing_score(t) < 1.0:
            continue
        legal = bool(C._NOT_TEXT.search(t))
        table = C.is_table_like(t)
        for q in QUESTIONS:
            if legal and not q.allow_legal:
                continue
            if not q.rx.search(t):
                continue
            if q.need_signal and not _VERDICT_CUE.search(t):
                continue
            if q.grupa == "Problemy" and table and q.id not in ("P04",):
                continue                                     # wiersze tabel jako „problem” tylko dla zaległości (tam to dane)
            rows.append(dict(doc_id=doc_id, miasto=city, typ_dokumentu=doc_type, grupa=q.grupa, pytanie_id=q.id, pytanie=q.nazwa, strona=u.page, id_cytatu=i,
                             cytat=t, typ="tabela/kwoty" if table else "zdanie", status=status_of(q.grupa, t, q.id),
                             slowa_kluczowe=", ".join(dict.fromkeys(m.group(0).lower() for m in q.rx.finditer(t)))[:100]))
    return rows


def city_card(rows: list[dict]) -> list[dict]:
    """Agreguje cytaty do odpowiedzi miasto × pytanie."""
    by = defaultdict(list)
    for r in rows:
        by[(r["miasto"], r["pytanie_id"])].append(r)
    out = []
    for q in QUESTIONS:
        for city in sorted({r["miasto"] for r in rows}):
            rs = by.get((city, q.id), [])
            cnt = Counter(r["status"] for r in rs)
            if rs:
                ans = "; ".join(f"{k} ×{v}" for k, v in cnt.most_common())
            else:
                ans = "nie znaleziono w dokumentach"
            seen, ex = set(), []
            for r in sorted(rs, key=lambda r: (r["typ"] != "zdanie", r["typ_dokumentu"] != "WPGMZ", r["strona"])):
                if r["id_cytatu"] not in seen and len(ex) < 3:
                    seen.add((r["doc_id"], r["id_cytatu"]))
                    seen.add(r["id_cytatu"])
                    ex.append(f"[{r['typ_dokumentu']}] s.{r['strona']} ({r['status']}): {r['cytat'][:230]}")
            out.append(dict(miasto=city, grupa=q.grupa, pytanie_id=q.id, pytanie=q.nazwa, odpowiedz=ans, n=len(rs),
                            kod="".join(f"{CODES[k]}{v} " for k, v in cnt.items()).strip(), dokumenty=", ".join(sorted({r["typ_dokumentu"] for r in rs})), przyklady="\n".join(ex)))
    return out


def append_sheets(xlsx_path: str | Path, rows: list[dict]) -> Path:
    """Dopisuje do istniejącego skoroszytu arkusze Karta_miasta (długi), Macierz (miasto × pytanie) i Cytaty_karta (wszystkie cytaty)."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = load_workbook(xlsx_path)
    card = city_card(rows)
    cities = sorted({c["miasto"] for c in card})
    FILL = {"Działania własne": "E2EFDA", "Zarządzanie zasobem": "DDEBF7", "Obowiązki ustawowe": "EDEDED", "Problemy": "FCE4D6", "Co udało się / ocena": "FFF2CC"}

    def style_header(ws):
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864"); c.alignment = Alignment(wrap_text=True, vertical="center")

    ws = wb.create_sheet("Karta_miasta")
    ws.append(["miasto", "grupa", "id", "pytanie", "odpowiedź (statusy cytatów)", "n cytatów", "dokumenty", "przykładowe cytaty (do 3)"])
    for c in card:
        ws.append([c["miasto"], c["grupa"], c["pytanie_id"], c["pytanie"], c["odpowiedz"], c["n"], c["dokumenty"], c["przyklady"]])
    style_header(ws)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10); cell.alignment = Alignment(wrap_text=True, vertical="top")
        if row[4].value == "nie znaleziono w dokumentach":
            row[4].font = Font(name="Arial", size=10, color="808080")
    for i, w in enumerate([16, 20, 6, 48, 34, 9, 12, 120], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions

    wm = wb.create_sheet("Macierz")
    wm.append(["miasto"] + [f"{q.id} {q.nazwa}" for q in QUESTIONS])
    code = {(c["miasto"], c["pytanie_id"]): c["kod"] for c in card}
    for city in cities:
        wm.append([city] + [code.get((city, q.id), "") for q in QUESTIONS])
    style_header(wm)
    for j, q in enumerate(QUESTIONS, 2):
        wm.cell(1, j).fill = PatternFill("solid", fgColor=FILL[q.grupa]); wm.cell(1, j).font = Font(name="Arial", bold=True, size=9)
        for i in range(2, len(cities) + 2):
            c = wm.cell(i, j); c.font = Font(name="Arial", size=10); c.alignment = Alignment(horizontal="center", vertical="center")
            if c.value:
                c.fill = PatternFill("solid", fgColor=FILL[q.grupa])
    wm.row_dimensions[1].height = 95
    for j in range(1, len(QUESTIONS) + 2):
        wm.column_dimensions[get_column_letter(j)].width = 18 if j == 1 else 13
        if j > 1:
            wm.cell(1, j).alignment = Alignment(wrap_text=True, vertical="top", text_rotation=0)
    wm.freeze_panes = "B2"

    wq = wb.create_sheet("Cytaty_karta")
    wq.append(["doc_id", "miasto", "typ dok.", "grupa", "id", "pytanie", "strona PDF", "status", "typ", "CYTAT", "słowa kluczowe", "id_cytatu"])
    for r in rows:
        wq.append([r["doc_id"], r["miasto"], r["typ_dokumentu"], r["grupa"], r["pytanie_id"], r["pytanie"], r["strona"], r["status"], r["typ"], r["cytat"], r["slowa_kluczowe"], r["id_cytatu"]])
    style_header(wq)
    for row in wq.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10); cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate([28, 14, 8, 20, 6, 40, 8, 18, 11, 110, 28, 9], 1):
        wq.column_dimensions[get_column_letter(i)].width = w
    wq.freeze_panes = "A2"; wq.auto_filter.ref = wq.dimensions

    leg = wb["Legenda"]
    for k, v in (("Karta miasta", "Arkusz Karta_miasta: odpowiedź na każde pytanie badawcze (budowa mieszkań komunalnych i socjalnych, TBS, SIM, grunty, wsparcie zakupu, preferencje dla młodych, problemy, obowiązki ustawowe, co udało się, plany) z cytatami. "
                                  "„nie znaleziono w dokumentach” ≠ „miasto tego nie robi”."),
                 ("Macierz", "Arkusz Macierz (miasto × pytanie): kody statusów i liczba cytatów, np. I3 P1 Z2 = 3 cytaty o istniejących/realizowanych działaniach, 1 planowany, 2 zrealizowane. "
                             "I = istniejące/realizowane, P = planowane, Z = zrealizowane, W = warunkowe, N = zaprzeczone/wstrzymane, D = wskazano jako problem (diagnoza). Kolory: zielony = działania własne, niebieski = zarządzanie, szary = obowiązki ustawowe, czerwony = problemy, żółty = wyniki/oceny."),
                 ("Plany i deklaracje", "Plany to cytaty ze statusem „planowane” (w kolumnie status): filtruj Cytaty_karta po status = planowane. Deklaracje ogólne bez nazwanego instrumentu są w arkuszu Własna_polityka (kolumny G).")):
        leg.append([k, v])
        for c in leg[leg.max_row]:
            c.font = Font(name="Arial", size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(xlsx_path)
    return Path(xlsx_path)
