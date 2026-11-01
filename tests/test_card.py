"""Testy karty miasta (bez modelu językowego)."""
import os, sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
import pytest

from hp import card as K, reader
from hp.extract_llm import fold_alnum
from hp.ingest import load_pages

HAVE = os.path.exists("data/interim/WPGMZ_Gdynia_2022-2026.pages.jsonl") and os.path.exists("data/interim/WPGMZ_Olsztyn_2022-2026.pages.jsonl")
need_pages = pytest.mark.skipif(not HAVE, reason="uruchom najpierw etap ingest")


def answers(text, doc_type="WPGMZ"):
    from hp.ingest import Page
    return K.answer_questions("d", doc_type, "X", [Page("d", 1, text, "text", len(text), len(text))])


def test_young_people_preference_and_sim_are_found_in_strategy_text():
    r = answers("Wdrażanie miejskiej polityki mieszkaniowej przewidującej preferencje w najmie mieszkań dla młodych rodzin z wykorzystaniem formuły SIM.", "SR")
    ids = {x["pytanie_id"] for x in r}
    assert "A09" in ids and "A03" in ids


def test_sales_halt_is_negated_and_past_results_are_marked_done():
    r = answers("Sprzedaż gminnych lokali mieszkalnych na rzecz najemców nie będzie prowadzona przez gminę w latach 2022-2026.")
    assert [x["status"] for x in r if x["pytanie_id"] == "B01"] == ["zaprzeczone/wstrzymane"]
    r = answers("Do końca roku 2021 Gmina zlikwidowała 201 pieców w gminnych budynkach mieszkalnych.")
    assert any(x["pytanie_id"] == "W01" and x["status"] == "zrealizowane" for x in r)


def test_problem_topics_cover_user_examples():
    text = ("W zasobie znajdują się pustostany wymagające remontu. Zaległości najemców w opłatach wzrosły. "
            "Niewykonanie obowiązku dostarczenia lokalu powoduje roszczenie o zapłatę odszkodowania. Brak mieszkań i deficyt lokali socjalnych.")
    ids = {x["pytanie_id"] for x in answers(text)}
    assert {"P01", "P03", "P04", "P05"} <= ids


@need_pages
def test_cards_for_gdynia_and_olsztyn_with_verbatim_quotes():
    rows = []
    for d, c in (("WPGMZ_Gdynia_2022-2026", "Gdynia"), ("WPGMZ_Olsztyn_2022-2026", "Olsztyn")):
        pages = load_pages(d, "data/interim")
        rr = K.answer_questions(d, "WPGMZ", c, pages)
        byp = {p.page: fold_alnum(p.text) for p in pages}
        assert rr and all(fold_alnum(r["cytat"]) in byp[r["strona"]] for r in rr)           # tylko dosłowne cytaty
        rows += rr
    card = {(c["miasto"], c["pytanie_id"]): c for c in K.city_card(rows)}
    assert card[("Gdynia", "A02")]["n"] > 0 and card[("Gdynia", "A10")]["n"] > 0              # TBS „Czynszówka”, mieszkania wspierane
    assert "zaprzeczone/wstrzymane" in card[("Gdynia", "B01")]["odpowiedz"]                   # wstrzymanie sprzedaży
    assert card[("Olsztyn", "A12")]["n"] > 0 and card[("Olsztyn", "A01")]["n"] > 0           # Fundusz Dopłat/BGK, budowa przy ul. Towarowej
    assert card[("Olsztyn", "P04")]["n"] >= 5                                                # zaległości — główny problem Olsztyna


@need_pages
def test_card_sheets_are_appended_to_workbook(tmp_path):
    from openpyxl import load_workbook
    pages = load_pages("WPGMZ_Gdynia_2022-2026", "data/interim")
    rows, s = reader.read_document("WPGMZ_Gdynia_2022-2026", "WPGMZ", pages, "Gdynia")
    p = reader.export_xlsx(rows, [s], tmp_path / "c.xlsx", {})
    K.append_sheets(p, K.answer_questions("WPGMZ_Gdynia_2022-2026", "WPGMZ", "Gdynia", pages))
    wb = load_workbook(p)
    assert {"Karta_miasta", "Macierz", "Cytaty_karta"} <= set(wb.sheetnames) and wb["Macierz"].max_column == len(K.QUESTIONS) + 1
