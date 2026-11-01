"""Testy czytnika cytatów (bez modelu językowego)."""
import os, sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
import pytest

from hp import candidates as C, reader
from hp.extract_llm import fold_alnum
from hp.ingest import load_pages

HAVE = os.path.exists("data/interim/WPGMZ_Gdynia_2022-2026.pages.jsonl")
need_pages = pytest.mark.skipif(not HAVE, reason="uruchom najpierw etap ingest")


def tags(text, doc_type="WPGMZ"):
    return reader.tag_unit(C.Unit(1, text), doc_type)


def test_sales_halt_is_instrument_I21_with_restrict_direction():
    r = [x for x in tags("Sprzedaż gminnych lokali mieszkalnych na rzecz najemców nie będzie prowadzona.") if x["kategoria"] == "Działanie"]
    assert r and r[0]["temat"].startswith("I21") and r[0]["kierunek"] == "ograniczanie/wstrzymanie"


def test_negated_action_does_not_count_for_own_policy():
    rows = [dict(r, doc_id="d", miasto="X", typ_dokumentu="WPGMZ", id_cytatu=1) for r in tags("Gmina nie planuje budowy mieszkań komunalnych w latach 2022-2026.")]
    assert any(x["status"] == "zaprzeczone" for x in rows if x["kategoria"] == "Działanie")
    assert reader.own_policy(rows)["A"]["n"] == 0


def test_swot_diagnosis_without_action_signal_does_not_make_city_own_policy():
    t = "Duża liczba oczekujących na mieszkania komunalne oraz spory zasób pustostanów wymagających remontu"
    rows = [dict(r, doc_id="d", miasto="X", typ_dokumentu="SR", id_cytatu=1) for r in tags(t, "SR")]
    assert reader.own_policy(rows)["A"]["n"] == 0


def test_table_rows_stay_in_finance_but_not_in_problems():
    rows = tags("remonty lokali mieszkalnych - pustostany 7 500 000 likwidacja pieców na opał stały 5 000 000 termomodernizacja budynków 5 600 000")
    assert not any(r["kategoria"] == "Problem" for r in rows)
    assert any(r["typ"] == "tabela/kwoty" for r in rows)


def test_legal_basis_and_definitions_are_skipped():
    assert tags("Stosownie do art. 21 ust. 1 ustawy z dnia 21 czerwca 2001 r. Rada Gminy uchwala wieloletnie programy gospodarowania mieszkaniowym zasobem") == []


@need_pages
def test_every_quote_is_verbatim_on_the_cited_page_and_gdynia_has_own_policy():
    pages = load_pages("WPGMZ_Gdynia_2022-2026", "data/interim")
    rows, s = reader.read_document("WPGMZ_Gdynia_2022-2026", "WPGMZ", pages, "Gdynia")
    by_page = {p.page: fold_alnum(p.text) for p in pages}
    assert rows and all(fold_alnum(r["cytat"]) in by_page[r["strona"]] for r in rows)       # żadnych wymyślonych fraz
    assert any(r["kategoria"] == "Problem" and "P02" in r["temat"] and "1945" in r["cytat"] for r in rows)
    assert any(r["kategoria"] == "Działanie" and r["temat"].startswith("I21") for r in rows)
    v = reader.city_verdict([s])
    assert v["werdykt"].startswith("TAK") and any("TBS" in e or "wspieran" in e for e in v["A"]["examples"])


@need_pages
def test_excel_export_has_expected_sheets(tmp_path):
    from openpyxl import load_workbook
    pages = load_pages("WPGMZ_Gdynia_2022-2026", "data/interim")
    rows, s = reader.read_document("WPGMZ_Gdynia_2022-2026", "WPGMZ", pages, "Gdynia")
    p = reader.export_xlsx(rows, [s], tmp_path / "c.xlsx", {"WPGMZ_Gdynia_2022-2026": dict(resolution_number="XXXIV/1127/21", adoption_date="2021-10-27")})
    wb = load_workbook(p)
    assert wb.sheetnames == ["Legenda", "Dokumenty", "Cytaty", "Własna_polityka"] and wb["Cytaty"].max_row == len(rows) + 1


def test_general_housing_declaration_in_strategy_is_marked_G_not_A():
    q = "Rozwój zrównoważonego budownictwa mieszkaniowego, w tym tworzenie i wdrażanie instrumentów wspierających rodziny w zaspokajaniu potrzeb mieszkaniowych"
    rows = [dict(r, doc_id="d", miasto="X", typ_dokumentu="SR", id_cytatu=1) for r in tags(q, "SR")]
    assert any(r["kategoria"] == "Działanie" and r["poziom"] == "G" for r in rows)
    s = reader.summarize("d", "SR", "X", rows, [])
    assert reader.city_verdict([s])["werdykt"].startswith("OGÓLNA DEKLARACJA")
