"""Generuje codebook/codebook.xlsx z modeli Pydantic i słowników (jedno źródło prawdy: schema.py + taxonomy.py)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import schema as S
from .taxonomy import CROSSWALK, FAMILIES, GROUPS, INSTRUMENTS, PROBLEMS, SPECIFICITY_RULE, WPGMZ_ELEMENTS

FONT = "Arial"
METADATA_VARS = [
    ("M01", "doc_id", "tekst", "Unikalny identyfikator dokumentu (typ_miasto_okres).", "rejestr"),
    ("M02", "city_id", "tekst", "Kod TERYT miasta na prawach powiatu (7 cyfr).", "rejestr"),
    ("M03", "doc_type", "kategoria", "SR = strategia rozwoju; WPGMZ = wieloletni program gospodarowania mieszkaniowym zasobem gminy.", "rejestr"),
    ("M04", "resolution_number", "tekst", "Numer uchwały rady miasta przyjmującej dokument (przy nowelizacji: numer uchwały nowelizującej).", "automatycznie + weryfikacja"),
    ("M05", "adoption_date", "data ISO", "Data uchwalenia. Daty nieistniejące w kalendarzu są flagowane, nie poprawiane.", "automatycznie + weryfikacja"),
    ("M06", "publication_date / publication_journal / publication_pos", "data / tekst / liczba", "Publikacja w dzienniku urzędowym województwa (jeśli jest w pliku).", "automatycznie"),
    ("M07", "period_start, period_end", "rok", "Pierwszy i ostatni rok obowiązywania/horyzontu dokumentu (dla „do 2030” period_start puste).", "automatycznie + weryfikacja"),
    ("M08", "is_amendment, amends_number, amends_date", "tak/nie, tekst, data", "Czy plik zawiera uchwałę zmieniającą wcześniejszą (np. Biała Podlaska 2022 zmienia uchwałę z 2021).", "automatycznie"),
    ("M09", "replaces_number", "tekst", "Uchwała, która traci moc w związku z przyjęciem dokumentu.", "automatycznie"),
    ("M10", "text_extraction", "kategoria", "text / ocr / mixed; liczba stron OCR; flagi jakości.", "automatycznie"),
]
EXAMPLE = ["WPGMZ_gdynia_2022-2026", "stock.units_total", "Gdynia", "5254", "5254 lokali gminnych", "2", "s. 1 „Aktualny zasób … tworzy 5 254 lokali gminnych”",
           "Liczba jak w dokumencie; wspólnoty wliczone do zasobu."]


def _rows_from_models():
    rows = []
    for bname, bmodel, elems in S.WPGMZ_BLOCKS:
        for sub, f in bmodel.model_fields.items():
            sm = f.annotation
            for var, vf in sm.model_fields.items():
                ex = vf.json_schema_extra or {}
                rows.append(("WPGMZ", sm.__name__, f"{sub}.{var}", ex.get("type") or ("lista" if "list" in str(vf.annotation) else "—"),
                             vf.description or "", ex.get("rq", ""), ", ".join(elems) or "cały dokument"))
    for var, vf in S.StrategyFrame.model_fields.items():
        ex = vf.json_schema_extra or {}
        rows.append(("SR", "StrategyFrame", f"frame.{var}", ex.get("type") or "lista", vf.description or "", ex.get("rq", ""), "pasaże o mieszkalnictwie"))
    return rows


def _items_rows():
    rows = []
    for cls, rq in ((S.ProblemItem, "RQ1, RQ4"), (S.GoalItem, "RQ2"), (S.ActionItem, "RQ2, RQ4, RQ5")):
        for var, vf in cls.model_fields.items():
            rows.append((cls.__name__, var, str(vf.annotation).replace("typing.", "").replace("<class '", "").replace("'>", ""), vf.description or "", rq))
    return rows


def _sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(list(r))
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT, size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def export(path: str | Path = "codebook/codebook.xlsx") -> Path:
    wb = Workbook(); wb.remove(wb.active)
    legend = [
        ("Księga kodowa v0.1", "Lokalna polityka mieszkaniowa w miastach na prawach powiatu: strategie rozwoju (SR) i WPGMZ."),
        ("Źródło prawdy", "Arkusz generowany z src/hp/schema.py i src/hp/taxonomy.py — nie edytuj ręcznie, zmieniaj kod i uruchom `python -m hp.cli codebook`."),
        ("Status wartości", "found = wartość wprost w dokumencie; not_present = dokument nie zawiera informacji (brak strukturalny); unclear = niejednoznaczne/sprzeczne."),
        ("Dowód", "Każda wartość found ma stronę i dosłowny cytat; cytaty są weryfikowane programowo."),
        ("Zasada porównywalności", "Zmienne W1 koduje się „jak w dokumencie”; normalizację (np. lokale na 1000 mieszkańców) robi aggregate.py z danymi BDL GUS."),
        ("Elementy ustawowe WPGMZ", "; ".join(f"{k} = {v}" for k, v in WPGMZ_ELEMENTS.items())),
        ("Arkusz „Arkusz_dwoch_kodujacych”", "Szablon do ręcznego podwójnego kodowania próby walidacyjnej; pierwszy wiersz to przykład wypełnienia."),
    ]
    _sheet(wb, "Legenda", ["Element", "Opis"], legend, [32, 120])
    _sheet(wb, "Metadane", ["ID", "Zmienna", "Typ", "Definicja", "Sposób pozyskania"], METADATA_VARS, [8, 40, 20, 90, 28])
    _sheet(wb, "Zmienne_dokumentow", ["Typ dok.", "Blok", "Zmienna", "Typ", "Definicja / instrukcja", "RQ", "Źródło w dokumencie"], _rows_from_models(), [10, 14, 40, 14, 100, 8, 26])
    _sheet(wb, "Pozycje_problem_cel_dzialanie", ["Tabela", "Pole", "Typ", "Definicja", "RQ"], _items_rows(), [16, 24, 40, 90, 14])
    _sheet(wb, "Problemy_P", ["Kod", "Nazwa", "Definicja", "Wyłączenia / granice"], [(k, v["name"], v["definition"], v["exclude"]) for k, v in PROBLEMS.items()], [8, 44, 90, 60])
    _sheet(wb, "Instrumenty_I", ["Kod", "Rodzina", "Nazwa", "Definicja"], [(k, f"{v['family']} {FAMILIES[v['family']]}", v["name"], v["definition"]) for k, v in INSTRUMENTS.items()], [8, 36, 56, 90])
    _sheet(wb, "Crosswalk_P_do_I", ["Problem", "Nazwa problemu", "Instrumenty „adresujące” (hipoteza robocza)"],
           [(k, PROBLEMS[k]["name"], ", ".join(v)) for k, v in CROSSWALK.items()], [10, 50, 80])
    _sheet(wb, "Skale", ["Skala", "Zasada"], [("Konkretność działania", SPECIFICITY_RULE),
           ("Nasilenie problemu", "mention = wzmianka; notable = opisany wyraźnie (≥2 zdania lub dane); key = wskazany jako wyzwanie kluczowe/priorytetowe"),
           ("Integracja mieszkalnictwa w strategii", "none → spatial_only → single_direction → dedicated_goal_or_priority → dedicated_programme_or_section"),
           ("Grupy adresatów", ", ".join(GROUPS))], [30, 120])
    _sheet(wb, "Arkusz_dwoch_kodujacych", ["doc_id", "zmienna", "miasto", "kodujący A", "kodujący B", "status", "strona", "uwagi"], [EXAMPLE], [26, 26, 14, 18, 18, 10, 40, 40])
    p = Path(path); p.parent.mkdir(parents=True, exist_ok=True); wb.save(p)
    return p


if __name__ == "__main__":
    print(export())
